"""Refresh the tariff-type table from the official Forum Datenaustausch source.

The list changes regularly (TARMED ended 2025-12-31; TARDOC and the ambulatory
flat rates began 2026-01-01), so this is one command:

    paperless-tarmed refresh-tariffs --dry-run    # show what changed
    paperless-tarmed refresh-tariffs              # write it

It scrapes the Referenzdaten page for the *current* download link rather than
hardcoding a URL, because the filename carries a publication date. Changes are
printed so a refresh is reviewable instead of an opaque file rewrite.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from .tariffs import BUNDLED_DATA_PATH, official_labels, writable_data_path

REFERENZDATEN_URL = "https://www.forum-datenaustausch.ch/referenzdaten"
BASE_URL = "https://www.forum-datenaustausch.ch"
# The page puts the real file URL in a print-only <span> as '[https://...xlsx]',
# while the <a href> is a wrapper route -- so match the URL anywhere, not in href.
LINK_RE = re.compile(r"""[^\s"'\[\]<>]*Tariftyp[^\s"'\[\]<>]*\.xlsx""", re.IGNORECASE)

# (sheet name, status to record). Terminated tariffs are kept: we process
# historical documents, and a 2024 invoice can legitimately cite a code that has
# since been retired.
SHEETS = [("Tariftyp-Liste", "aktiv"), ("Terminiert Tariftyp", "terminiert")]

COL_CODE, COL_STATUS, COL_DE, COL_FR = 0, 1, 2, 3
COL_VALID_FROM, COL_VALID_TO = 37, 38


def find_xlsx_url(client) -> str:
    resp = client.get(REFERENZDATEN_URL, follow_redirects=True, timeout=60)
    resp.raise_for_status()
    matches = LINK_RE.findall(resp.text)
    if not matches:
        raise SystemExit(
            f"No 'Tariftyp*.xlsx' link found on {REFERENZDATEN_URL}.\n"
            "The page layout may have changed -- open it and check by hand."
        )
    url = matches[0]
    return url if url.startswith("http") else f"{BASE_URL}{url}"


def parse_workbook(blob: bytes) -> dict[str, dict]:
    import io

    try:
        import openpyxl
    except ModuleNotFoundError:  # pragma: no cover - depends on install extras
        raise SystemExit(
            "openpyxl is required to parse the tariff workbook.\n"
            "Install it with:  pip install openpyxl"
        ) from None

    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    entries: dict[str, dict] = {}
    for sheet_name, default_status in SHEETS:
        if sheet_name not in wb.sheetnames:
            print(f"  ! sheet {sheet_name!r} missing, skipped", file=sys.stderr)
            continue
        for row in wb[sheet_name].iter_rows(min_row=5, values_only=True):
            if not row or not row[COL_CODE]:
                continue
            code = str(row[COL_CODE]).strip()
            # Skip header repeats and range placeholders ('H00-H99').
            if not re.fullmatch(r"[0-9A-Z]{3}", code):
                continue

            def cell(idx: int, _row=row) -> str | None:
                if idx >= len(_row) or _row[idx] is None:
                    return None
                text = str(_row[idx]).strip()
                return text[:10] if idx in (COL_VALID_FROM, COL_VALID_TO) else text

            entries[code] = {
                "code": code,
                "status": cell(COL_STATUS) or default_status,
                "label_fr": cell(COL_FR),
                "label_de": cell(COL_DE),
                "valid_from": cell(COL_VALID_FROM),
                "valid_to": cell(COL_VALID_TO),
            }
    return entries


def diff(old: dict[str, dict], new: dict[str, dict]) -> list[str]:
    lines = []
    for code in sorted(set(new) - set(old)):
        lines.append(f"  + {code}  {new[code]['label_fr']}")
    for code in sorted(set(old) - set(new)):
        lines.append(f"  - {code}  {old[code].get('label_fr')}")
    for code in sorted(set(old) & set(new)):
        if old[code] != new[code]:
            was, now = old[code].get("status"), new[code].get("status")
            detail = f"{was} -> {now}" if was != now else "details updated"
            lines.append(f"  ~ {code}  {new[code]['label_fr']}  ({detail})")
    return lines


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        prog="paperless-tarmed refresh-tariffs",
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    ap.add_argument("--url", help="Bypass scraping and download this XLSX directly.")
    ap.add_argument("--out", type=Path, help="Where to write (default: the active table).")
    ap.add_argument("--dry-run", action="store_true", help="Report changes without writing.")
    args = ap.parse_args(argv)

    import httpx

    with httpx.Client(headers={"User-Agent": "paperless-tarmed"}) as client:
        url = args.url or find_xlsx_url(client)
        print(f"source: {url}")
        resp = client.get(url, follow_redirects=True, timeout=120)
        resp.raise_for_status()
        entries = parse_workbook(resp.content)

    if not entries:
        raise SystemExit("Parsed zero tariff types -- refusing to overwrite the table.")

    old = official_labels()
    changes = diff(old, entries)
    active = sum(1 for e in entries.values() if e["status"] == "aktiv")
    print(f"parsed {len(entries)} tariff types ({active} active)")
    if old:
        print("\n".join(changes) if changes else "  (no changes)")

    if args.dry_run:
        print("dry run -- nothing written")
        return 0

    out = args.out or writable_data_path()
    try:
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(
            json.dumps(
                {"source_url": url, "tariff_types": [entries[c] for c in sorted(entries)]},
                ensure_ascii=False,
                indent=2,
            )
            + "\n",
            encoding="utf-8",
        )
    except OSError as exc:
        raise SystemExit(
            f"Cannot write {out}: {exc}\n"
            "Inside the container, set PAPERLESS_TARMED_TARIFF_PATH to a writable "
            "location on the data volume (the image sets /data/tariff_types.json)."
        ) from exc
    official_labels.cache_clear()
    if out != BUNDLED_DATA_PATH:
        print(f"wrote {out} (overrides the table bundled in the image)")
    else:
        print(f"wrote {out}")
    return 0
